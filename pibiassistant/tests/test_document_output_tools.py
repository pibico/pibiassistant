"""get_document_pdf, export_data and send_email with attached PDFs on real documents."""

import io
import unittest
from unittest import mock

import frappe

from pibiassistant.pibiassistant_chat.api.chat.aida_tools import READ_TOOLS, WRITE_TOOLS, _approval_card, chat_tool_specs
from pibiassistant.plugins.core.tools.export_data import DataExport
from pibiassistant.plugins.core.tools.get_document_pdf import DocumentPdf
from pibiassistant.plugins.pao.tools.send_email import SendEmail


class TestOutputTools(unittest.TestCase):
    def setUp(self):
        frappe.set_user("Administrator")
        self.invoice = frappe.db.get_value("Sales Invoice", {"docstatus": 1}, "name")
        if not self.invoice:
            self.skipTest("the site has no submitted Sales Invoice")
        self.files = []
        self.todos = []

    def tearDown(self):
        frappe.set_user("Administrator")
        for url in self.files:
            for name in frappe.get_all("File", filters={"file_url": url}, pluck="name"):
                frappe.delete_doc("File", name, force=True, ignore_permissions=True)
        for name in self.todos:
            frappe.delete_doc("ToDo", name, force=True, ignore_permissions=True)
        frappe.db.commit()

    def _content(self, url):
        name = frappe.db.get_value("File", {"file_url": url}, "name")
        content = frappe.get_doc("File", name).get_content()
        return content if isinstance(content, bytes) else content.encode("utf-8")

    def test_pdf_of_a_document_is_private_and_real(self):
        result = DocumentPdf().execute({"doctype": "Sales Invoice", "name": self.invoice})
        self.assertTrue(result["success"], result)
        self.files.append(result["file_url"])
        self.assertTrue(result["file_url"].startswith("/private/files/"))
        self.assertTrue(self._content(result["file_url"]).startswith(b"%PDF"))
        self.assertIn(result["file_url"], result["download_link"])
        again = DocumentPdf().execute({"doctype": "Sales Invoice", "name": self.invoice})
        self.files.append(again["file_url"])
        self.assertEqual(frappe.db.count("File", {"file_name": result["file_name"], "is_private": 1}), 1)

    def test_pdf_refuses_bad_input_and_guests(self):
        bad = DocumentPdf().execute({"doctype": "Sales Invoice", "name": self.invoice, "print_format": "No Such Format"})
        self.assertFalse(bad["success"])
        self.assertIn("Available", bad["error"])
        self.assertFalse(DocumentPdf().execute({"doctype": "Sales Invoice", "name": "ZZ-none"})["success"])
        frappe.set_user("Guest")
        try:
            self.assertFalse(DocumentPdf().execute({"doctype": "Sales Invoice", "name": self.invoice})["success"])
        finally:
            frappe.set_user("Administrator")

    def _export(self, **args):
        result = DataExport().execute({"doctype": "Customer", **args})
        if result.get("file_url"):
            self.files.append(result["file_url"])
        return result

    def test_export_csv_and_xlsx(self):
        csv_result = self._export(format="csv", fields=["name", "customer_name"], limit=5)
        self.assertTrue(csv_result["success"], csv_result)
        self.assertEqual(csv_result["columns"], ["name", "customer_name"])
        data = self._content(csv_result["file_url"]).decode("utf-8-sig").splitlines()
        self.assertEqual(data[0], "name,customer_name")
        self.assertEqual(len(data) - 1, csv_result["rows"])

        xlsx_result = self._export(format="xlsx", fields=["name"], limit=3)
        self.assertTrue(xlsx_result["success"], xlsx_result)
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(self._content(xlsx_result["file_url"]))).active
        self.assertEqual(sheet.cell(1, 1).value, "name")
        self.assertEqual(sheet.max_row - 1, xlsx_result["rows"])

    def test_export_limit_truncates_and_says_so(self):
        result = self._export(format="csv", fields=["name"], limit=2)
        self.assertEqual(result["rows"], 2)
        self.assertTrue(result["truncated"])

    def test_export_never_includes_secrets_or_credential_doctypes(self):
        self.assertIn("cannot be exported", DataExport().execute({"doctype": "User", "fields": ["api_secret"]})["error"])
        self.assertIn("cannot be exported", DataExport().execute({"doctype": "OAuth Client"})["error"])
        self.assertIn("Unknown fields", DataExport().execute({"doctype": "Customer", "fields": ["no_such_field"]})["error"])
        self.assertFalse(DataExport().execute({"doctype": "Customer", "format": "pdf"})["success"])
        self.assertFalse(DataExport().execute({"doctype": "Customer", "order_by": "name; drop table x"})["success"])

    def test_export_does_not_let_a_cell_become_a_formula(self):
        todo = frappe.get_doc({"doctype": "ToDo", "description": "=1+1 ZZ-EXPORT-TEST"}).insert(ignore_permissions=True)
        self.todos.append(todo.name)
        frappe.db.commit()
        result = DataExport().execute({"doctype": "ToDo", "fields": ["name", "description"], "format": "csv", "filters": {"name": todo.name}})
        self.files.append(result["file_url"])
        self.assertIn("'=1+1 ZZ-EXPORT-TEST", self._content(result["file_url"]).decode("utf-8-sig"))

    def test_export_applies_filter_repairs_and_guests_are_refused(self):
        result = self._export(format="csv", fields=["name"], filters={"creation": ["between", "2000-01-01", "2100-01-01"]}, limit=2)
        self.assertTrue(result["success"], result)
        self.assertEqual(result["rows"], 2)
        frappe.set_user("Guest")
        try:
            self.assertFalse(DataExport().execute({"doctype": "Customer", "fields": ["name"]})["success"])
        finally:
            frappe.set_user("Administrator")

    def test_email_attaches_the_pdf_of_a_document(self):
        with mock.patch.object(frappe, "sendmail") as sendmail:
            result = SendEmail().execute(
                {"recipients": ["destino@example.com"], "subject": "Factura", "message": "Adjunta",
                 "attach_documents": [{"doctype": "Sales Invoice", "name": self.invoice}]}
            )
        self.assertTrue(result["success"], result)
        kwargs = sendmail.call_args.kwargs
        self.assertEqual(len(kwargs["attachments"]), 1)
        self.assertTrue(kwargs["attachments"][0]["fcontent"].startswith(b"%PDF"))
        self.assertTrue(kwargs["attachments"][0]["fname"].endswith(".pdf"))
        self.assertEqual(result["attachments"], [kwargs["attachments"][0]["fname"]])

    def test_email_attachment_failures_send_nothing(self):
        with mock.patch.object(frappe, "sendmail") as sendmail:
            for docs in (
                [{"doctype": "Sales Invoice", "name": "ZZ-none"}],
                [{"doctype": "Sales Invoice", "name": self.invoice}] * 4,
                ["not a dict"],
            ):
                result = SendEmail().execute({"recipients": ["destino@example.com"], "subject": "x", "message": "y", "attach_documents": docs})
                self.assertFalse(result["success"], docs)
            sendmail.assert_not_called()

    def test_email_without_attachments_is_unchanged(self):
        with mock.patch.object(frappe, "sendmail") as sendmail:
            result = SendEmail().execute({"recipients": ["destino@example.com"], "subject": "x", "message": "y"})
        self.assertTrue(result["success"])
        self.assertIsNone(sendmail.call_args.kwargs["attachments"])

    def test_chat_policy_and_card_previews(self):
        self.assertIn("get_document_pdf", READ_TOOLS)
        frappe.set_user("Administrator")
        specs = {s["function"]["name"] for s in chat_tool_specs("Administrator")}
        self.assertIn("get_document_pdf", specs)
        self.assertNotIn("get_document_pdf", WRITE_TOOLS)
        self.assertIn("export_data", WRITE_TOOLS)
        card = _approval_card({"interrupt_id": "i1", "name": "export_data", "arguments": {"doctype": "Customer", "format": "csv"}})
        self.assertIn("Export Customer to csv", card["reason"]["description"])
        card = _approval_card(
            {"interrupt_id": "i2", "name": "send_email",
             "arguments": {"recipients": ["a@example.com"], "subject": "Factura", "message": "m",
                           "attach_documents": [{"doctype": "Sales Invoice", "name": self.invoice}]}}
        )
        self.assertIn(f"Attaches the PDF of Sales Invoice {self.invoice}", card["reason"]["description"])
